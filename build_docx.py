#!/usr/bin/env python3
"""Compile the venue-chatbot design docs (01-12) into a single Word .docx."""
import re
import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = "/Users/wcities/ReactStudy/chatbot"
OUT = os.path.join(SRC, "Venue-Chatbot-Platform-Research.docx")

FILES = [
    "01-vision.md",
    "02-what-works-vs-what-doesnt.md",
    "03-architecture.md",
    "04-data-model.md",
    "05-chat-orchestrator.md",
    "06-multi-vertical-modules.md",
    "07-integrations.md",
    "08-safety-and-guardrails.md",
    "09-llm-and-tech-stack.md",
    "10-roadmap.md",
    "11-team-budget-risks.md",
    "12-decisions-now-vs-later.md",
]

ACCENT = RGBColor(0x1F, 0x4E, 0x79)   # deep blue
GREY = RGBColor(0x55, 0x55, 0x55)
CODE_BG = "F2F2F2"

# ---- money-stripping rules (commercial/finance content removed; product
# features that merely involve currency, e.g. menu prices & refund guardrails,
# are kept) ----
SECTION_SKIP = {
    "09-llm-and-tech-stack.md": ["cost per venue", "infra cost estimate"],
    "11-team-budget-risks.md": ["budget"],
}

LINE_DROP = [
    "$40k+ MRR", "$20k+ MRR", "$30k+ MRR",
    "$40k MRR, ~$500k ARR run-rate",
    "Series A pitch ready",
    "Pricing (approximate, 2026):",
    "$3 per 1M input tokens", "$15 per 1M output tokens",
    "$2.50 per 1M input tokens", "$10 per 1M output tokens",
    "**Pricing model**",
]

REPLACE = [
    ("US-based, $200k–$5M annual revenue, owner-operated, 1–20 staff.",
     "US-based, owner-operated, 1–20 staff."),
    ("and will pay\n$50–200/mo for software that saves them hours.",
     "and want software that saves them hours."),
    ("If these are hit, the Series A story is credible. ", ""),
    ("| A custom-trained LLM | $10M+ compute, beaten by frontier",
     "| A custom-trained LLM | Beaten by frontier"),
    ("tour operators, $40k+/month revenue, ready for\nSeries A.",
     "tour operators."),
    ("Restaurant economics validated. Hotels = higher ACV ($200–500/mo vs\n"
     "$50–100 for restaurants), justifying the engineering investment.",
     "Restaurant model validated. Hotels are the next vertical, justifying "
     "the engineering investment."),
    ("Frontier LLM training costs $10M–$100M+ in compute, 10–20 ML researchers,",
     "Frontier LLM training needs 10–20 ML researchers,"),
    ("| Train your own model | $10M+, 12+ months, immediately behind frontier |",
     "| Train your own model | 12+ months, immediately behind frontier |"),
    ("catches runaway venues\n   before they cost $500.",
     "catches runaway venues early."),
    ("Quarterly third-party pen test once profitable ($5k–15k)",
     "Quarterly third-party pen test"),
    ("Cyber liability insurance ($1–5M policy)", "Cyber liability insurance"),
    ("- **SOC 2 Type I** (then Type II) — start within 12 months. $40–80k Year 1.",
     "- **SOC 2 Type I** (then Type II) — start within 12 months."),
    ("$10M+ compute, 10–20 researchers,", "10–20 researchers,"),
    ("pick lowest for $1M coverage", "pick lowest for adequate coverage"),
    ("Templates if <$1M raised; counsel once revenue justifies",
     "Templates early; counsel later"),
    ("$1–3k for a designer, before customer #11",
     "Engage a designer before customer #11"),
    (" (one venue runs up $500 in a day)", ""),
    ("alerts at $50/day/venue", "alerts on anomalies"),
]


def strip_money(fname, md):
    skips = SECTION_SKIP.get(fname, [])
    if skips:
        out = []
        skip_level = None
        for line in md.split("\n"):
            hm = re.match(r"^(#{1,6})\s+(.*)$", line)
            if hm:
                level = len(hm.group(1))
                title = hm.group(2).strip().lower()
                if skip_level is not None and level <= skip_level:
                    skip_level = None
                if skip_level is None and any(s in title for s in skips):
                    skip_level = level
                    continue
            if skip_level is not None:
                continue
            out.append(line)
        md = "\n".join(out)
    for old, new in REPLACE:
        md = md.replace(old, new)
    md = "\n".join(l for l in md.split("\n")
                   if not any(d in l for d in LINE_DROP))
    return md


doc = Document()

# ---- base styles ----
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15

for lvl, sz, color in [(1, 18, ACCENT), (2, 14, ACCENT), (3, 12, ACCENT)]:
    st = doc.styles[f"Heading {lvl}"]
    st.font.name = "Calibri"
    st.font.size = Pt(sz)
    st.font.color.rgb = color
    st.font.bold = True


def shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def add_inline(paragraph, text):
    """Parse **bold**, *italic*, `code`, [text](url) into runs."""
    token = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`|\[[^\]]+?\]\([^)]+?\))")
    pos = 0
    for m in token.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        t = m.group(0)
        if t.startswith("**"):
            r = paragraph.add_run(t[2:-2]); r.bold = True
        elif t.startswith("*"):
            r = paragraph.add_run(t[1:-1]); r.italic = True
        elif t.startswith("`"):
            r = paragraph.add_run(t[1:-1])
            r.font.name = "Consolas"; r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)
        elif t.startswith("["):
            lm = re.match(r"\[([^\]]+)\]\(([^)]+)\)", t)
            r = paragraph.add_run(lm.group(1))
            r.font.color.rgb = ACCENT; r.underline = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def add_code_block(lines):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Inches(0.2)
    pf.space_before = Pt(4); pf.space_after = Pt(8)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), CODE_BG)
    pPr.append(shd)
    r = p.add_run("\n".join(lines))
    r.font.name = "Consolas"; r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)


def add_table(header, rows):
    cols = len(header)
    t = doc.add_table(rows=1, cols=cols)
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, h in enumerate(header):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        add_inline(p, h.strip())
        for run in p.runs:
            run.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade_cell(hdr[i], "1F4E79")
    for row in rows:
        cells = t.add_row().cells
        for i in range(cols):
            cells[i].text = ""
            val = row[i].strip() if i < len(row) else ""
            add_inline(cells[i].paragraphs[0], val)
    doc.add_paragraph()


def parse_table_row(line):
    line = line.strip().strip("|")
    return [c.strip() for c in line.split("|")]


def render_markdown(md):
    lines = md.split("\n")
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()

        # code fence
        if stripped.startswith("```"):
            block = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                block.append(lines[i]); i += 1
            add_code_block(block)
            i += 1
            continue

        # table
        if "|" in line and i + 1 < n and re.match(r"^\s*\|?[\s:|-]+\|?\s*$", lines[i + 1]) and "-" in lines[i + 1]:
            header = parse_table_row(line)
            i += 2
            rows = []
            while i < n and "|" in lines[i] and lines[i].strip():
                rows.append(parse_table_row(lines[i])); i += 1
            add_table(header, rows)
            continue

        # blank
        if not stripped:
            i += 1
            continue

        # horizontal rule
        if re.match(r"^\s*(-{3,}|\*{3,}|_{3,})\s*$", line):
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            pbdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:space"), "1"); bottom.set(qn("w:color"), "CCCCCC")
            pbdr.append(bottom); pPr.append(pbdr)
            i += 1
            continue

        # headings
        hm = re.match(r"^(#{1,6})\s+(.*)$", line)
        if hm:
            level = len(hm.group(1))
            text = hm.group(2).strip()
            p = doc.add_heading(level=min(level, 4))
            add_inline(p, text)
            i += 1
            continue

        # blockquote
        if stripped.startswith(">"):
            quote_lines = []
            while i < n and lines[i].strip().startswith(">"):
                quote_lines.append(lines[i].strip().lstrip(">").strip()); i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            pPr = p._p.get_or_add_pPr()
            pbdr = OxmlElement("w:pBdr")
            left = OxmlElement("w:left")
            left.set(qn("w:val"), "single"); left.set(qn("w:sz"), "18")
            left.set(qn("w:space"), "8"); left.set(qn("w:color"), "1F4E79")
            pbdr.append(left); pPr.append(pbdr)
            add_inline(p, " ".join(quote_lines))
            for r in p.runs:
                r.italic = True; r.font.color.rgb = GREY
            continue

        # unordered list
        um = re.match(r"^(\s*)[-*+]\s+(.*)$", line)
        if um:
            indent = len(um.group(1))
            p = doc.add_paragraph(style="List Bullet")
            if indent >= 2:
                p.paragraph_format.left_indent = Inches(0.5 + 0.25 * (indent // 2))
            add_inline(p, um.group(2).strip())
            i += 1
            continue

        # ordered list
        om = re.match(r"^(\s*)\d+\.\s+(.*)$", line)
        if om:
            p = doc.add_paragraph(style="List Number")
            add_inline(p, om.group(2).strip())
            i += 1
            continue

        # paragraph
        p = doc.add_paragraph()
        add_inline(p, stripped)
        i += 1


# ---------- Title page ----------
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
t.paragraph_format.space_before = Pt(120)
r = t.add_run("Venue Chatbot Platform")
r.font.size = Pt(30); r.font.bold = True; r.font.color.rgb = ACCENT

st = doc.add_paragraph()
st.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = st.add_run("Product Research & Technical Design")
r.font.size = Pt(16); r.font.color.rgb = GREY

st2 = doc.add_paragraph()
st2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = st2.add_run(
    "A chat-controlled, multi-vertical management system for small venues —\n"
    "restaurants, hotels, tour operators, and retail."
)
r.font.size = Pt(11); r.italic = True; r.font.color.rgb = GREY

d = doc.add_paragraph()
d.alignment = WD_ALIGN_PARAGRAPH.CENTER
d.paragraph_format.space_before = Pt(180)
r = d.add_run("Prepared for: Client Review\nDate: June 2026\nStatus: Pre-build — Design Phase")
r.font.size = Pt(11); r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

doc.add_page_break()

# ---------- Contents ----------
doc.add_heading("Contents", level=1)
titles = {
    "01-vision.md": "1. Vision",
    "02-what-works-vs-what-doesnt.md": "2. What Works vs. What Doesn't",
    "03-architecture.md": "3. Architecture",
    "04-data-model.md": "4. Data Model",
    "05-chat-orchestrator.md": "5. Chat Orchestrator",
    "06-multi-vertical-modules.md": "6. Multi-Vertical Modules",
    "07-integrations.md": "7. Integrations",
    "08-safety-and-guardrails.md": "8. Safety & Guardrails",
    "09-llm-and-tech-stack.md": "9. LLM & Tech Stack",
    "10-roadmap.md": "10. Roadmap",
    "11-team-budget-risks.md": "11. Team & Risks",
    "12-decisions-now-vs-later.md": "12. Decisions: Now vs. Later",
}
for f in FILES:
    p = doc.add_paragraph(style="List Bullet")
    add_inline(p, titles[f])
doc.add_page_break()

# ---------- Body ----------
for idx, f in enumerate(FILES):
    with open(os.path.join(SRC, f), encoding="utf-8") as fh:
        md = fh.read()
    md = strip_money(f, md)
    # Drop the first H1 from each file; replace with our numbered section title
    md_lines = md.split("\n")
    if md_lines and md_lines[0].startswith("# "):
        md_lines = md_lines[1:]
    doc.add_heading(titles[f], level=1)
    render_markdown("\n".join(md_lines))
    if idx < len(FILES) - 1:
        doc.add_page_break()

# ---------- Footer with page numbers ----------
section = doc.sections[0]
footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = fp.add_run("Venue Chatbot Platform — Confidential   |   Page ")
run.font.size = Pt(8); run.font.color.rgb = GREY
fldChar1 = OxmlElement("w:fldChar"); fldChar1.set(qn("w:fldCharType"), "begin")
instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = "PAGE"
fldChar2 = OxmlElement("w:fldChar"); fldChar2.set(qn("w:fldCharType"), "end")
run2 = fp.add_run(); run2.font.size = Pt(8)
run2._r.append(fldChar1); run2._r.append(instr); run2._r.append(fldChar2)

doc.save(OUT)
print("Saved:", OUT)
